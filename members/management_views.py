from collections import defaultdict
from io import BytesIO

from django.conf import settings
from django.contrib import messages
from django.contrib.admin.views.decorators import staff_member_required
from django.core.paginator import Paginator
from django import forms
from django.db.models import Count, Q
from django.http import FileResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from openpyxl import load_workbook

from audit.models import AuditLog
from config.choices import Status
from families.models import Household
from families.models import ClanIdentity
from families.identity import clan_name, clan_prefix

from .forms import BulkMemberImportForm, PersonForm, RelationshipForm
from .models import CorrectionRequest, Marriage, Person, Relationship
from .services.bulk_import import import_members_from_workbook
from .services.family_tree import build_family_tree
from .services.relationships import synchronize_parent_relationship
from .views import LIST_PAGE_SIZE, query_without


class ClanIdentityForm(forms.ModelForm):
    class Meta:
        model = ClanIdentity
        fields = ("clan_name",)
        labels = {"clan_name": "Clan name (Jina la ukoo)"}

    def clean_clan_name(self):
        return self.cleaned_data["clan_name"].strip()


@staff_member_required
def clan_settings(request):
    identity, _ = ClanIdentity.objects.get_or_create(pk=1)
    form = ClanIdentityForm(request.POST or None, instance=identity)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Clan name has been updated across the registry.")
        return redirect("management-clan-settings")
    return render(request, "management/clan_settings.html", {"form": form})


@staff_member_required
def command_center(request):
    stats = {
        "pending_members": Person.objects.filter(status__in=pending_statuses()).count(),
        "pending_relationships": Relationship.objects.filter(status__in=pending_statuses()).count(),
        "pending_corrections": CorrectionRequest.objects.filter(status__in=pending_statuses()).count(),
        "households": Household.objects.count(),
    }
    recent_members = Person.objects.order_by("-created_at")[:8]
    return render(request, "management/command_center.html", {"stats": stats, "recent_members": recent_members})


@staff_member_required
def member_management(request):
    query = request.GET.get("q", "").strip()
    members = Person.objects.select_related("created_by").order_by("-created_at")
    if query:
        members = members.filter(
            Q(member_id__icontains=query)
            | Q(first_name__icontains=query)
            | Q(last_name__icontains=query)
            | Q(phone_number__icontains=query)
            | Q(email__icontains=query)
        )
    return render(
        request,
        "management/member_management.html",
        {
            "members": Paginator(members, LIST_PAGE_SIZE).get_page(request.GET.get("page")),
            "query": query,
            "bulk_import_form": BulkMemberImportForm(),
            "pagination_query": query_without(request, "page"),
        },
    )


@staff_member_required
def member_import_template(request):
    template_path = settings.BASE_DIR / "static" / "downloads" / "member_bulk_import_template.xlsx"
    return FileResponse(
        template_path.open("rb"),
        as_attachment=True,
        filename=f"{clan_prefix().lower()}_member_bulk_import_template.xlsx",
    )


@staff_member_required
def member_import_sample(request):
    sample_path = settings.BASE_DIR / "static" / "downloads" / "member_bulk_sample_data.xlsx"
    workbook = load_workbook(sample_path)
    sheet = workbook["Members"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    for row in sheet.iter_rows(min_row=2):
        clan_cell = row[headers["Clan name"] - 1]
        if clan_cell.value == "Moshi":
            clan_cell.value = clan_name()
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return FileResponse(
        output,
        as_attachment=True,
        filename=f"{clan_prefix().lower()}_member_bulk_sample_data.xlsx",
    )


@staff_member_required
def member_bulk_import(request):
    if request.method != "POST":
        return redirect("management-members")

    form = BulkMemberImportForm(request.POST, request.FILES)
    if not form.is_valid():
        members = Person.objects.select_related("created_by").order_by("-created_at")
        return render(
            request,
            "management/member_management.html",
            {
                "members": Paginator(members, LIST_PAGE_SIZE).get_page(1),
                "query": "",
                "bulk_import_form": form,
                "pagination_query": "",
            },
            status=400,
        )

    imported, import_errors = import_members_from_workbook(form.cleaned_data["workbook"], request.user)
    if import_errors:
        members = Person.objects.select_related("created_by").order_by("-created_at")
        return render(
            request,
            "management/member_management.html",
            {
                "members": Paginator(members, LIST_PAGE_SIZE).get_page(1),
                "query": "",
                "bulk_import_form": form,
                "import_errors": import_errors,
                "pagination_query": "",
            },
            status=400,
        )

    messages.success(request, f"{len(imported)} members were imported successfully.")
    return redirect("management-members")


@staff_member_required
def member_create(request):
    if request.method == "POST":
        form = PersonForm(request.POST, request.FILES)
        if form.is_valid():
            person = form.save(commit=False)
            person.created_by = request.user
            if not person.status:
                person.status = Status.VERIFIED
            person.save()
            messages.success(request, f"{person.full_name} has been added.")
            return redirect("management-members")
    else:
        form = PersonForm(initial={"status": Status.VERIFIED})
    return render(request, "management/member_form.html", {"form": form, "title": "Add Member"})


@staff_member_required
def member_update(request, pk):
    person = get_object_or_404(Person, pk=pk)
    if request.method == "POST":
        form = PersonForm(request.POST, request.FILES, instance=person)
        if form.is_valid():
            form.save()
            messages.success(request, f"{person.full_name} has been updated.")
            return redirect("management-members")
    else:
        form = PersonForm(instance=person)
    return render(request, "management/member_form.html", {"form": form, "title": "Edit Member", "person": person})


@staff_member_required
def member_delete(request, pk):
    person = get_object_or_404(Person, pk=pk)
    if request.method == "POST":
        name = person.full_name
        person.delete()
        messages.success(request, f"{name} has been deleted.")
        return redirect("management-members")
    return render(request, "management/confirm_delete.html", {"object": person, "cancel_url": "management-members"})


@staff_member_required
def relationship_management(request):
    relationships = Relationship.objects.select_related("from_person", "to_person", "verified_by").order_by("-created_at")
    return render(
        request,
        "management/relationship_management.html",
        {
            "relationships": Paginator(relationships, LIST_PAGE_SIZE).get_page(request.GET.get("page")),
            "pagination_query": query_without(request, "page"),
        },
    )


@staff_member_required
def relationship_create(request):
    if request.method == "POST":
        form = RelationshipForm(request.POST)
        if form.is_valid():
            relationship = form.save(commit=False)
            if relationship.status == Status.VERIFIED:
                relationship.verified_by = request.user
                relationship.verified_at = timezone.now()
            relationship.save()
            if relationship.status == Status.VERIFIED:
                try:
                    synchronize_parent_relationship(relationship)
                except ValueError as error:
                    relationship.delete()
                    form.add_error(None, str(error))
                    return render(request, "management/relationship_form.html", {"form": form, "title": "Add Relationship"})
                create_reciprocal_relationship(relationship, request.user)
            messages.success(request, "Relationship has been added.")
            return redirect("management-relationships")
    else:
        form = RelationshipForm(initial={"status": Status.VERIFIED})
    return render(request, "management/relationship_form.html", {"form": form, "title": "Add Relationship"})


@staff_member_required
def relationship_update(request, pk):
    relationship = get_object_or_404(Relationship, pk=pk)
    if request.method == "POST":
        form = RelationshipForm(request.POST, instance=relationship)
        if form.is_valid():
            relationship = form.save(commit=False)
            if relationship.status == Status.VERIFIED and not relationship.verified_by:
                relationship.verified_by = request.user
                relationship.verified_at = timezone.now()
            relationship.save()
            if relationship.status == Status.VERIFIED:
                try:
                    synchronize_parent_relationship(relationship)
                except ValueError as error:
                    form.add_error(None, str(error))
                    return render(request, "management/relationship_form.html", {"form": form, "title": "Edit Relationship", "relationship": relationship})
                create_reciprocal_relationship(relationship, request.user)
            messages.success(request, "Relationship has been updated.")
            return redirect("management-relationships")
    else:
        form = RelationshipForm(instance=relationship)
    return render(request, "management/relationship_form.html", {"form": form, "title": "Edit Relationship", "relationship": relationship})


@staff_member_required
def relationship_delete(request, pk):
    relationship = get_object_or_404(Relationship, pk=pk)
    if request.method == "POST":
        relationship.delete()
        messages.success(request, "Relationship has been deleted.")
        return redirect("management-relationships")
    return render(request, "management/confirm_delete.html", {"object": relationship, "cancel_url": "management-relationships"})


@staff_member_required
def pending_verification(request):
    if request.method == "POST":
        model_name = request.POST.get("model")
        object_id = request.POST.get("object_id")
        action = request.POST.get("action")
        status = Status.VERIFIED if action == "approve" else Status.REJECTED
        obj = get_pending_object(model_name, object_id)
        if isinstance(obj, Relationship) and status == Status.VERIFIED:
            try:
                synchronize_parent_relationship(obj)
            except ValueError as error:
                messages.error(request, str(error))
                return redirect("management-pending")
        obj.status = status
        if hasattr(obj, "verified_by") and status == Status.VERIFIED:
            obj.verified_by = request.user
            obj.verified_at = timezone.now()
        if hasattr(obj, "reviewed_by"):
            obj.reviewed_by = request.user
            obj.reviewed_at = timezone.now()
        obj.save()
        if isinstance(obj, Relationship) and status == Status.VERIFIED:
            create_reciprocal_relationship(obj, request.user)
        AuditLog.objects.create(
            actor=request.user,
            action=action,
            model_name=obj.__class__.__name__,
            object_id=str(obj.pk),
            field_name="status",
            new_value=status,
            approver=request.user,
        )
        messages.success(request, f"{obj} was {status.replace('_', ' ')}.")
        return redirect("management-pending")

    context = {
        "members": Paginator(Person.objects.filter(status__in=pending_statuses()).order_by("-created_at"), LIST_PAGE_SIZE).get_page(request.GET.get("members_page")),
        "relationships": Paginator(Relationship.objects.filter(status__in=pending_statuses()).select_related("from_person", "to_person").order_by("-created_at"), LIST_PAGE_SIZE).get_page(request.GET.get("relationships_page")),
        "marriages": Paginator(Marriage.objects.filter(status__in=pending_statuses()).select_related("spouse_one", "spouse_two").order_by("-pk"), LIST_PAGE_SIZE).get_page(request.GET.get("marriages_page")),
        "corrections": Paginator(CorrectionRequest.objects.filter(status__in=pending_statuses()).select_related("person", "requested_by").order_by("-created_at"), LIST_PAGE_SIZE).get_page(request.GET.get("corrections_page")),
        "members_query": query_without(request, "members_page"),
        "relationships_query": query_without(request, "relationships_page"),
        "marriages_query": query_without(request, "marriages_page"),
        "corrections_query": query_without(request, "corrections_page"),
    }
    return render(request, "management/pending_verification.html", context)


@staff_member_required
def duplicate_review(request):
    people = Person.objects.order_by("last_name", "first_name", "date_of_birth")
    groups = defaultdict(list)
    for person in people:
        key = (person.first_name.lower(), person.last_name.lower(), person.date_of_birth)
        if person.first_name and person.last_name:
            groups[key].append(person)
    duplicate_groups = Paginator([group for group in groups.values() if len(group) > 1], LIST_PAGE_SIZE).get_page(
        request.GET.get("groups_page")
    )
    phone_duplicates = (
        Person.objects.exclude(phone_number="")
        .values("phone_number")
        .annotate(total=Count("id"))
        .filter(total__gt=1)
        .order_by("phone_number")
    )
    return render(
        request,
        "management/duplicate_review.html",
        {
            "duplicate_groups": duplicate_groups,
            "phone_duplicates": Paginator(phone_duplicates, LIST_PAGE_SIZE).get_page(request.GET.get("phones_page")),
            "groups_query": query_without(request, "groups_page"),
            "phones_query": query_without(request, "phones_page"),
        },
    )


@staff_member_required
def tree_management(request):
    root_id = request.GET.get("root")
    root = Person.objects.filter(pk=root_id).first() if root_id else Person.objects.order_by("date_of_birth", "id").first()
    roots = Person.objects.order_by("last_name", "first_name")
    relationships = Relationship.objects.none()
    tree = {"nodes": [], "edges": []}
    if root:
        relationships = Relationship.objects.filter(
            Q(from_person=root) | Q(to_person=root)
        ).select_related("from_person", "to_person", "verified_by").order_by("-created_at")
        tree = build_family_tree(root, depth=2)
    return render(
        request,
        "management/tree_management.html",
        {
            "root": root,
            "roots": roots,
            "tree": tree,
            "relationships": Paginator(relationships, LIST_PAGE_SIZE).get_page(request.GET.get("page")),
            "pagination_query": query_without(request, "page"),
        },
    )


def pending_statuses():
    return [Status.DRAFT, Status.SUBMITTED, Status.PENDING_VERIFICATION]


def get_pending_object(model_name, object_id):
    model_map = {
        "person": Person,
        "relationship": Relationship,
        "marriage": Marriage,
        "correction": CorrectionRequest,
    }
    model = model_map.get(model_name)
    if model is None:
        raise ValueError("Unsupported verification model.")
    return get_object_or_404(model, pk=object_id)


def create_reciprocal_relationship(relationship, user):
    reciprocal_type = relationship.reciprocal_type
    if not reciprocal_type:
        return None
    reciprocal, created = Relationship.objects.get_or_create(
        from_person=relationship.to_person,
        to_person=relationship.from_person,
        relationship_type=reciprocal_type,
        defaults={
            "status": Status.VERIFIED,
            "verified_by": user,
            "verified_at": timezone.now(),
            "notes": "Auto-created when reciprocal relationship was approved.",
        },
    )
    if not created and reciprocal.status != Status.VERIFIED:
        reciprocal.status = Status.VERIFIED
        reciprocal.verified_by = user
        reciprocal.verified_at = timezone.now()
        reciprocal.save(update_fields=["status", "verified_by", "verified_at"])
    return reciprocal if created else None
