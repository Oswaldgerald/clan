from datetime import date, datetime

from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from config.choices import Status
from families.identity import clan_name

from ..forms import OUTSIDE_TANZANIA, PersonForm
from ..models import Person


MEMBER_IMPORT_HEADERS = (
    "Member ID",
    "First name *",
    "Middle name",
    "Last name *",
    "Clan name",
    "Gender",
    "Date of birth",
    "Place of birth",
    "Nationality",
    "Current residence",
    "Country",
    "Occupation",
    "Education",
    "Phone number",
    "Email",
    "Is living",
    "Date of death",
    "Place of death",
    "Burial location",
    "Family branch",
    "Generation",
    "Relationship to founder",
    "Status",
)


def import_members_from_workbook(uploaded_file, created_by):
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError) as exc:
        return [], [f"The uploaded file is not a readable Excel workbook: {exc}"]

    if "Members" not in workbook.sheetnames:
        workbook.close()
        return [], ["The workbook must contain a sheet named 'Members'."]

    sheet = workbook["Members"]
    headers = tuple(cell.value for cell in sheet[1][: len(MEMBER_IMPORT_HEADERS)])
    if headers != MEMBER_IMPORT_HEADERS:
        workbook.close()
        return [], ["The Members sheet headings were changed. Download a fresh template and keep the headings unchanged."]

    pending_people = []
    errors = []
    seen_member_ids = set()
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=2, max_col=len(MEMBER_IMPORT_HEADERS), values_only=True),
        start=2,
    ):
        if not any(value not in (None, "") for value in row):
            continue
        if row_number > 1001:
            errors.append("The workbook can contain at most 1,000 member rows.")
            break

        values = dict(zip(MEMBER_IMPORT_HEADERS, row))
        member_id = _text(values["Member ID"])
        if member_id and member_id in seen_member_ids:
            errors.append(f"Row {row_number}: Member ID '{member_id}' is repeated in the workbook.")
            continue
        if member_id:
            seen_member_ids.add(member_id)

        gender = _text(values["Gender"]).lower()
        if gender not in {Person.Gender.MALE, Person.Gender.FEMALE}:
            errors.append(f"Row {row_number}: Gender must be male or female.")
            continue

        living_value = _text(values["Is living"]).lower()
        if living_value not in {"", "yes", "no"}:
            errors.append(f"Row {row_number}: Is living must be yes or no.")
            continue

        status = _text(values["Status"]).lower() or Status.VERIFIED
        if status not in Status.values:
            errors.append(f"Row {row_number}: Status '{status}' is not valid.")
            continue

        country = _text(values["Country"])
        residence = normalize_residence(_text(values["Current residence"]), country)
        place_of_birth = normalize_residence(_text(values["Place of birth"]), country)
        data = {
            "member_id": member_id,
            "first_name": _text(values["First name *"]),
            "middle_name": _text(values["Middle name"]),
            "last_name": _text(values["Last name *"]),
            "clan_name": _text(values["Clan name"]) or clan_name(),
            "gender": gender,
            "date_of_birth": _date_text(values["Date of birth"]),
            "place_of_birth": place_of_birth,
            "current_residence": residence,
            "country": country,
            "occupation": _text(values["Occupation"]),
            "education": _text(values["Education"]),
            "phone_number": _text(values["Phone number"]),
            "email": _text(values["Email"]),
            "is_living": "on" if living_value != "no" else "",
            "date_of_death": _date_text(values["Date of death"]),
            "place_of_death": _text(values["Place of death"]),
            "burial_location": _text(values["Burial location"]),
            "relationship_to_founder": _text(values["Relationship to founder"]),
            "status": status,
        }
        form = PersonForm(data=data, show_family_metadata=True)
        if not form.is_valid():
            details = "; ".join(
                f"{field}: {', '.join(messages)}" for field, messages in form.errors.items()
            )
            errors.append(f"Row {row_number}: {details}")
            continue

        person = form.save(commit=False)
        person.created_by = created_by
        pending_people.append(person)

    workbook.close()
    if errors:
        return [], errors

    with transaction.atomic():
        for person in pending_people:
            person.save()
    return pending_people, []


def _text(value):
    return "" if value is None else str(value).strip()


def normalize_residence(value, country):
    if country and country.casefold() not in {"tanzania", "united republic of tanzania"}:
        return OUTSIDE_TANZANIA
    district_to_region = {
        "mamba": "Kilimanjaro",
        "marangu": "Kilimanjaro",
        "moshi": "Kilimanjaro",
        "mwika": "Kilimanjaro",
        "rombo": "Kilimanjaro",
    }
    return district_to_region.get(value.casefold(), value)


def _date_text(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return _text(value)


def _integer_text(value):
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)
