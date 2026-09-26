from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = ("Member ID", "Name", "Gender", "Date of Birth", "Age", "Residence")


def build_active_member_workbook(members):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Active Members"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:F{max(1, len(members) + 1)}"

    header_fill = PatternFill("solid", fgColor="3B82C4")
    for column, heading in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=heading)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")

    for row, member in enumerate(members, start=2):
        values = (
            member.member_id,
            member.full_name,
            member.get_gender_display() or "Not recorded",
            member.date_of_birth,
            member.report_age,
            member.current_residence or "Not recorded",
        )
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row, column=column, value=value)
        if member.date_of_birth:
            sheet.cell(row=row, column=4).number_format = "dd mmm yyyy"

    widths = (22, 32, 16, 18, 10, 28)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 24

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output
