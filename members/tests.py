import tempfile
from datetime import timedelta
from io import BytesIO

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from families.models import FamilyBranch

from config.choices import Status

from .forms import MemberRelativeForm, PersonForm, RelationshipForm, generate_member_id
from families.models import ClanIdentity
from .models import Marriage, Person, Relationship
from .services.family_tree import build_family_tree


class DashboardAccessTests(TestCase):
    def test_anonymous_visitor_is_redirected_to_login(self):
        response = self.client.get(reverse("dashboard"))

        self.assertRedirects(response, f"{reverse('login')}?next={reverse('dashboard')}")

    def test_authenticated_member_can_open_dashboard(self):
        User.objects.create_user(username="dashboard-member", password="pass12345")
        self.client.login(username="dashboard-member", password="pass12345")

        response = self.client.get(reverse("dashboard"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Clan Dashboard")


class ClanIdentityTests(TestCase):
    def test_staff_can_change_clan_name_across_brand_and_new_records(self):
        User.objects.create_user(username="clan-admin", password="pass12345", is_staff=True)
        self.client.login(username="clan-admin", password="pass12345")

        response = self.client.post(reverse("management-clan-settings"), {"clan_name": "Mandara"})

        self.assertRedirects(response, reverse("management-clan-settings"))
        self.assertEqual(ClanIdentity.objects.get(pk=1).clan_name, "Mandara")
        self.assertContains(self.client.get(reverse("dashboard")), "Mandara Clan Registry")
        self.assertEqual(Person.objects.create(member_id="NEW-1", first_name="New", last_name="Member").clan_name, "Mandara")
        self.assertTrue(generate_member_id().startswith("MANDARA-"))

    def test_non_staff_cannot_change_clan_name(self):
        User.objects.create_user(username="clan-member", password="pass12345")
        self.client.login(username="clan-member", password="pass12345")

        response = self.client.post(reverse("management-clan-settings"), {"clan_name": "Other"})

        self.assertNotEqual(response.status_code, 200)
        self.assertFalse(ClanIdentity.objects.exists())


class FamilyTreeTests(TestCase):
    def setUp(self):
        self.branch = FamilyBranch.objects.create(name="Main Branch")
        self.father = Person.objects.create(
            member_id="MOSHI-001",
            first_name="Amani",
            last_name="Moshi",
            gender=Person.Gender.MALE,
            family_branch=self.branch,
            generation=1,
        )
        self.mother = Person.objects.create(
            member_id="MOSHI-002",
            first_name="Neema",
            last_name="Moshi",
            gender=Person.Gender.FEMALE,
            family_branch=self.branch,
            generation=1,
        )
        self.child = Person.objects.create(
            member_id="MOSHI-003",
            first_name="Baraka",
            last_name="Moshi",
            gender=Person.Gender.MALE,
            father=self.father,
            mother=self.mother,
            family_branch=self.branch,
            generation=2,
        )
        Marriage.objects.create(spouse_one=self.father, spouse_two=self.mother)
        Relationship.objects.create(
            from_person=self.father,
            to_person=self.child,
            relationship_type=Relationship.Type.FATHER,
        )

    def test_family_tree_service_returns_nodes_and_edges(self):
        tree = build_family_tree(self.child, depth=2)

        self.assertEqual(tree["root"], self.child.pk)
        self.assertEqual(len(tree["nodes"]), 3)
        self.assertIn(
            {"source": self.father.pk, "target": self.child.pk, "relation": "father"},
            tree["edges"],
        )

    def test_family_tree_data_endpoint(self):
        response = self.client.get(reverse("family-tree-data"), {"root": self.child.pk})

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["root"], self.child.pk)
        self.assertEqual(len(payload["nodes"]), 3)

    def test_family_tree_excludes_unverified_relationships(self):
        outsider = Person.objects.create(member_id="MOSHI-PENDING", first_name="Pending", last_name="Relative")
        Relationship.objects.create(
            from_person=self.child,
            to_person=outsider,
            relationship_type=Relationship.Type.SIBLING,
            status=Status.SUBMITTED,
        )

        tree = build_family_tree(self.child, depth=2)

        self.assertNotIn(outsider.pk, {node["id"] for node in tree["nodes"]})

    def test_family_tree_uses_default_depth_for_invalid_input(self):
        tree = build_family_tree(self.child, depth="invalid")

        self.assertEqual(tree["root"], self.child.pk)

    def test_marriage_does_not_make_spouse_an_automatic_parent(self):
        spouse = Person.objects.create(
            member_id="MOSHI-SPOUSE", first_name="Spouse", last_name="Moshi", gender=Person.Gender.FEMALE
        )
        child = Person.objects.create(
            member_id="MOSHI-OUTSIDE", first_name="Independent", last_name="Moshi", father=self.father
        )
        Marriage.objects.create(
            spouse_one=self.father,
            spouse_two=spouse,
            status=Status.VERIFIED,
        )

        tree = build_family_tree(child, depth=2)

        self.assertIn({"source": self.father.pk, "target": child.pk, "relation": "father"}, tree["edges"])
        self.assertNotIn({"source": spouse.pk, "target": child.pk, "relation": "mother"}, tree["edges"])


class PersonDetailRelationshipTests(TestCase):
    def test_parent_field_without_relationship_is_shown_once(self):
        father = Person.objects.create(
            member_id="MOSHI-FIELD-FATHER", first_name="Field", last_name="Father", gender=Person.Gender.MALE
        )
        child = Person.objects.create(
            member_id="MOSHI-FIELD-CHILD", first_name="Field", last_name="Child", father=father
        )

        response = self.client.get(reverse("person-detail", args=[father.pk]))

        self.assertEqual(response.context["children"].paginator.count, 1)
        self.assertEqual(response.context["relationships"].paginator.count, 1)
        self.assertContains(response, "Father of")
        self.assertEqual(Relationship.objects.count(), 0)

        Relationship.objects.create(
            from_person=father,
            to_person=child,
            relationship_type=Relationship.Type.FATHER,
            status=Status.VERIFIED,
        )
        response = self.client.get(reverse("person-detail", args=[father.pk]))
        self.assertEqual(response.context["relationships"].paginator.count, 1)

    def test_verified_father_relationship_is_listed_as_child(self):
        father = Person.objects.create(
            member_id="MOSHI-DETAIL-FATHER", first_name="Father", last_name="Moshi", gender=Person.Gender.MALE
        )
        child = Person.objects.create(
            member_id="MOSHI-DETAIL-CHILD", first_name="Child", last_name="Moshi", gender=Person.Gender.MALE
        )
        Relationship.objects.create(
            from_person=father,
            to_person=child,
            relationship_type=Relationship.Type.FATHER,
            status=Status.VERIFIED,
        )

        response = self.client.get(reverse("person-detail", args=[father.pk]))

        self.assertContains(response, child.full_name)
        self.assertContains(response, "Father of")

    def test_unverified_relationship_is_not_shown_on_public_profile(self):
        father = Person.objects.create(member_id="MOSHI-PENDING-FATHER", first_name="Pending", last_name="Father")
        child = Person.objects.create(member_id="MOSHI-PENDING-CHILD", first_name="Pending", last_name="Child")
        Relationship.objects.create(
            from_person=father,
            to_person=child,
            relationship_type=Relationship.Type.FATHER,
            status=Status.SUBMITTED,
        )

        response = self.client.get(reverse("person-detail", args=[father.pk]))

        self.assertNotContains(response, "Father of")
        self.assertContains(response, "No children recorded.")


class AddChildTests(TestCase):
    def setUp(self):
        User.objects.create_user(username="child-entry", password="pass12345")
        self.client.login(username="child-entry", password="pass12345")
        self.father = Person.objects.create(
            member_id="MOSHI-ADD-FATHER", first_name="Entry", last_name="Father", gender=Person.Gender.MALE
        )
        self.child = Person.objects.create(
            member_id="MOSHI-ADD-CHILD", first_name="Entry", last_name="Child"
        )

    def test_existing_child_is_submitted_once(self):
        url = reverse("add-child", args=[self.father.pk])
        response = self.client.get(url)
        self.assertContains(response, "Existing member")
        self.assertContains(response, "New member")
        self.assertContains(response, f"?parent={self.father.pk}")

        self.client.post(url, {"child": self.child.pk})
        self.client.post(url, {"child": self.child.pk})
        self.assertEqual(
            Relationship.objects.filter(
                from_person=self.father,
                to_person=self.child,
                relationship_type=Relationship.Type.FATHER,
                status=Status.SUBMITTED,
            ).count(),
            1,
        )
        self.child.refresh_from_db()
        self.assertIsNone(self.child.father_id)

    def test_existing_recorded_child_does_not_create_request(self):
        self.child.father = self.father
        self.child.save(update_fields=["father"])
        self.client.post(reverse("add-child", args=[self.father.pk]), {"child": self.child.pk})
        self.assertFalse(Relationship.objects.exists())

    def test_conflicting_father_is_rejected(self):
        another = Person.objects.create(
            member_id="MOSHI-OTHER-FATHER", first_name="Other", last_name="Father", gender=Person.Gender.MALE
        )
        self.child.father = another
        self.child.save(update_fields=["father"])
        response = self.client.post(reverse("add-child", args=[self.father.pk]), {"child": self.child.pk})
        self.assertContains(response, "already has a different father")
        self.assertFalse(Relationship.objects.exists())


class MemberListFilterTests(TestCase):
    def setUp(self):
        self.branch = FamilyBranch.objects.create(name="Northern Branch")
        self.male = Person.objects.create(
            member_id="MOSHI-LIST-1",
            first_name="Male",
            last_name="Member",
            gender=Person.Gender.MALE,
            family_branch=self.branch,
        )
        self.female = Person.objects.create(
            member_id="MOSHI-LIST-2",
            first_name="Female",
            last_name="Member",
            gender=Person.Gender.FEMALE,
            is_living=False,
        )

    def test_member_list_shows_gender_and_living_summaries(self):
        response = self.client.get(reverse("member-list"))

        self.assertEqual(response.context["stats"]["total"], 2)
        self.assertEqual(response.context["stats"]["male"], 1)
        self.assertEqual(response.context["stats"]["female"], 1)
        self.assertEqual(response.context["stats"]["deceased"], 1)

    def test_member_list_defaults_to_living_members(self):
        response = self.client.get(reverse("member-list"))

        self.assertEqual(response.context["living_filter"], "living")
        self.assertContains(response, self.male.full_name)
        self.assertNotContains(response, self.female.full_name)

    def test_member_list_can_explicitly_show_all_life_statuses(self):
        response = self.client.get(reverse("member-list"), {"living_status": ""})

        self.assertContains(response, self.male.full_name)
        self.assertContains(response, self.female.full_name)

    def test_member_list_combines_gender_living_and_branch_filters(self):
        response = self.client.get(
            reverse("member-list"),
            {"gender": "male", "living_status": "living", "branch": self.branch.pk},
        )

        self.assertContains(response, self.male.full_name)
        self.assertNotContains(response, self.female.full_name)

    def test_member_list_is_paginated_and_preserves_filters(self):
        Person.objects.bulk_create(
            [
                Person(
                    member_id=f"MOSHI-PAGE-{index}",
                    first_name=f"Paged{index}",
                    last_name="Member",
                    gender=Person.Gender.MALE,
                )
                for index in range(21)
            ]
        )

        response = self.client.get(reverse("member-list"), {"gender": "male"})

        self.assertEqual(response.context["members"].paginator.per_page, 10)
        self.assertEqual(response.context["members"].paginator.num_pages, 3)
        self.assertEqual(response.context["pagination_query"], "gender=male")
        self.assertContains(response, "Showing 1-10 of 22")


class PortalAccessTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_superuser(username="admin", email="admin@example.com", password="pass12345")
        self.member = User.objects.create_user(username="member", email="member@example.com", password="pass12345")

    def test_member_login_page_is_public(self):
        response = self.client.get(reverse("login"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Member Login")

    def test_member_profile_requires_login(self):
        response = self.client.get(reverse("member-profile"))

        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("login"), response["Location"])

    def test_account_settings_requires_login(self):
        response = self.client.get(reverse("account-settings"))

        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("login"), response["Location"])

    def test_member_profile_for_logged_in_member(self):
        self.client.login(username="member", password="pass12345")

        response = self.client.get(reverse("member-profile"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "My Profile")
        self.assertContains(response, "data-profile-image")
        self.assertContains(response, reverse("account-settings"))
        self.assertContains(response, "data-account-menu-button")

    def test_member_can_update_account_settings(self):
        self.client.login(username="member", password="pass12345")

        response = self.client.post(
            reverse("account-settings"),
            {
                "first_name": "Updated",
                "last_name": "Member",
                "email": "updated@example.com",
                "phone_number": "0712345678",
            },
        )

        self.assertRedirects(response, reverse("account-settings"))
        self.member.refresh_from_db()
        self.assertEqual(self.member.first_name, "Updated")
        self.assertEqual(self.member.phone_number, "0712345678")

    def test_member_can_upload_profile_picture_without_linked_person(self):
        self.client.login(username="member", password="pass12345")
        image = SimpleUploadedFile(
            "avatar.gif",
            b"GIF87a\x01\x00\x01\x00\x80\x00\x00\x00\x00\x00\xff\xff\xff!\xf9\x04\x01\x00\x00\x00\x00,\x00\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02D\x01\x00;",
            content_type="image/gif",
        )

        with tempfile.TemporaryDirectory() as media_root, self.settings(MEDIA_ROOT=media_root):
            response = self.client.post(
                reverse("member-profile"),
                {"update_photo": "1", "profile_photo": image},
            )

        self.assertRedirects(response, reverse("member-profile"))
        self.member.refresh_from_db()
        self.assertTrue(self.member.profile_photo.name.startswith("accounts/photos/avatar"))

    def test_linked_member_photo_is_shown_for_logged_in_user(self):
        profile = self.member.person_profile
        profile.member_id = "MOSHI-AVATAR"
        profile.first_name = "Photo"
        profile.last_name = "Member"
        profile.profile_photo = "members/photos/member.jpg"
        profile.save()
        self.client.login(username="member", password="pass12345")

        response = self.client.get(reverse("member-profile"))

        self.assertContains(response, "/media/members/photos/member.jpg")
        self.assertContains(response, "Profile photo of Photo Member")

    def test_every_account_is_available_as_a_relationship_person(self):
        self.assertEqual(self.staff.person_profile.account, self.staff)
        self.assertEqual(self.member.person_profile.account, self.member)

        form = RelationshipForm()

        self.assertIn(self.staff.person_profile, form.fields["from_person"].queryset)
        self.assertIn(self.member.person_profile, form.fields["to_person"].queryset)

    def test_member_list_includes_accounts_regardless_of_role(self):
        response = self.client.get(reverse("member-list"))

        self.assertContains(response, self.staff.person_profile.full_name)
        self.assertContains(response, self.member.person_profile.full_name)
        self.assertContains(response, self.staff.get_role_display())

    def test_management_requires_staff(self):
        self.client.login(username="member", password="pass12345")

        response = self.client.get(reverse("management-dashboard"))

        self.assertEqual(response.status_code, 302)

    def test_management_dashboard_for_staff(self):
        self.client.login(username="admin", password="pass12345")

        response = self.client.get(reverse("management-dashboard"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Management Command Center")

    def test_admin_management_alias_for_staff(self):
        self.client.login(username="admin", password="pass12345")

        response = self.client.get(reverse("admin-management-dashboard"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Management Command Center")

    def test_staff_can_create_member_from_management_page(self):
        self.client.login(username="admin", password="pass12345")

        response = self.client.post(
            reverse("management-member-add"),
            {
                "member_id": "",
                "first_name": "Juma",
                "middle_name": "",
                "last_name": "Moshi",
                "clan_name": "Moshi",
                "gender": Person.Gender.MALE,
                "is_living": "on",
                "status": Status.VERIFIED,
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertTrue(Person.objects.filter(first_name="Juma", status=Status.VERIFIED).exists())

    def test_staff_can_download_bulk_member_template(self):
        self.client.login(username="admin", password="pass12345")

        response = self.client.get(reverse("management-member-import-template"))

        self.assertEqual(response.status_code, 200)
        self.assertIn("moshi_member_bulk_import_template.xlsx", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(b"".join(response.streaming_content)), read_only=True)
        self.assertIn("Members", workbook.sheetnames)
        self.assertEqual(workbook["Members"]["B1"].value, "First name *")
        workbook.close()

    def test_staff_can_bulk_import_members_from_completed_template(self):
        self.client.login(username="admin", password="pass12345")
        template_path = settings.BASE_DIR / "static" / "downloads" / "member_bulk_import_template.xlsx"
        workbook = load_workbook(template_path)
        sheet = workbook["Members"]
        sheet["B2"] = "Bulk"
        sheet["D2"] = "Member"
        sheet["F2"] = "female"
        sheet["P2"] = "yes"
        sheet["W2"] = "verified"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        upload = SimpleUploadedFile(
            "completed_members.xlsx",
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(
            reverse("management-member-bulk-import"),
            {"workbook": upload},
        )

        self.assertRedirects(response, reverse("management-members"))
        self.assertTrue(
            Person.objects.filter(first_name="Bulk", last_name="Member", gender=Person.Gender.FEMALE).exists()
        )

    def test_generated_sample_workbook_is_upload_ready(self):
        self.client.login(username="admin", password="pass12345")
        members_before = Person.objects.count()
        sample_path = settings.BASE_DIR / "static" / "downloads" / "member_bulk_sample_data.xlsx"
        upload = SimpleUploadedFile(
            "member_bulk_sample_data.xlsx",
            sample_path.read_bytes(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(
            reverse("management-member-bulk-import"),
            {"workbook": upload},
        )

        self.assertRedirects(response, reverse("management-members"))
        self.assertEqual(Person.objects.count() - members_before, 15)

    def test_management_member_form_is_sectioned(self):
        self.client.login(username="admin", password="pass12345")

        response = self.client.get(reverse("management-member-add"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Identity (Utambulisho)")
        self.assertContains(response, "Demographics (Taarifa binafsi)")
        self.assertContains(response, "Location (Mahali)")
        self.assertContains(response, "Family Links (Mahusiano ya familia)")
        self.assertContains(response, "First name (Jina la kwanza)")
        self.assertContains(response, "Namba ya mwanaukoo")
        self.assertNotContains(response, "mwanachama")

    def test_member_can_submit_relative_for_approval(self):
        self.client.login(username="member", password="pass12345")

        response = self.client.post(
            reverse("submit-relative"),
            {
                "member_id": "",
                "first_name": "Sara",
                "middle_name": "",
                "last_name": "Moshi",
                "clan_name": "Moshi",
                "gender": Person.Gender.FEMALE,
                "is_living": "on",
            },
        )

        self.assertEqual(response.status_code, 302)
        person = Person.objects.get(first_name="Sara")
        self.assertEqual(person.status, Status.SUBMITTED)
        self.assertEqual(person.created_by, self.member)

    def test_member_relative_form_is_sectioned(self):
        self.client.login(username="member", password="pass12345")

        response = self.client.get(reverse("submit-relative"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Identity (Utambulisho)")
        self.assertContains(response, "Demographics (Taarifa binafsi)")
        self.assertContains(response, "Living Status (Hali ya uhai)")
        self.assertContains(response, "Father (Baba)")

    def test_relationship_form_has_swahili_labels(self):
        self.client.login(username="member", password="pass12345")

        response = self.client.get(reverse("submit-relationship"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Relationship type (Aina ya uhusiano)")
        self.assertContains(response, "Notes (Maelezo)")

    def test_member_can_submit_relationship_for_approval(self):
        first = Person.objects.create(member_id="MOSHI-010", first_name="One", last_name="Moshi")
        second = Person.objects.create(member_id="MOSHI-011", first_name="Two", last_name="Moshi")
        self.client.login(username="member", password="pass12345")

        response = self.client.post(
            reverse("submit-relationship"),
            {
                "from_person": first.pk,
                "relationship_type": Relationship.Type.SIBLING,
                "to_person": second.pk,
                "notes": "Submitted by member",
            },
        )

        self.assertEqual(response.status_code, 302)
        relationship = Relationship.objects.get(from_person=first, to_person=second)
        self.assertEqual(relationship.status, Status.SUBMITTED)

    def test_staff_can_approve_pending_relationship_and_create_reciprocal(self):
        first = Person.objects.create(member_id="MOSHI-020", first_name="Alpha", last_name="Moshi")
        second = Person.objects.create(member_id="MOSHI-021", first_name="Beta", last_name="Moshi")
        relationship = Relationship.objects.create(
            from_person=first,
            relationship_type=Relationship.Type.SIBLING,
            to_person=second,
            status=Status.SUBMITTED,
        )
        self.client.login(username="admin", password="pass12345")

        response = self.client.post(
            reverse("management-pending"),
            {"model": "relationship", "object_id": relationship.pk, "action": "approve"},
        )

        self.assertEqual(response.status_code, 302)
        relationship.refresh_from_db()
        self.assertEqual(relationship.status, Status.VERIFIED)
        self.assertTrue(
            Relationship.objects.filter(
                from_person=second,
                to_person=first,
                relationship_type=Relationship.Type.SIBLING,
                status=Status.VERIFIED,
            ).exists()
        )

    def test_approving_father_relationship_updates_child_parent(self):
        father = Person.objects.create(
            member_id="MOSHI-FATHER-1", first_name="Baba", last_name="Moshi", gender=Person.Gender.MALE
        )
        child = Person.objects.create(
            member_id="MOSHI-CHILD-1", first_name="Mtoto", last_name="Moshi", gender=Person.Gender.FEMALE
        )
        relationship = Relationship.objects.create(
            from_person=father,
            relationship_type=Relationship.Type.FATHER,
            to_person=child,
            status=Status.SUBMITTED,
        )
        self.client.login(username="admin", password="pass12345")

        response = self.client.post(
            reverse("management-pending"),
            {"model": "relationship", "object_id": relationship.pk, "action": "approve"},
        )

        self.assertEqual(response.status_code, 302)
        child.refresh_from_db()
        self.assertEqual(child.father, father)

    def test_pending_verification_page_supports_marriages_without_created_at(self):
        first = Person.objects.create(member_id="MOSHI-MARRIAGE-1", first_name="First", last_name="Moshi")
        second = Person.objects.create(member_id="MOSHI-MARRIAGE-2", first_name="Second", last_name="Moshi")
        marriage = Marriage.objects.create(
            spouse_one=first,
            spouse_two=second,
            status=Status.SUBMITTED,
        )
        self.client.login(username="admin", password="pass12345")

        response = self.client.get(reverse("management-pending"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, str(marriage))


class ParentGenderControlTests(TestCase):
    def setUp(self):
        self.male = Person.objects.create(
            member_id="MOSHI-100",
            first_name="Male",
            last_name="Parent",
            gender=Person.Gender.MALE,
        )
        self.female = Person.objects.create(
            member_id="MOSHI-101",
            first_name="Female",
            last_name="Parent",
            gender=Person.Gender.FEMALE,
        )
        self.unknown = Person.objects.create(
            member_id="MOSHI-102",
            first_name="Unknown",
            last_name="Parent",
            gender="",
        )

    def test_parent_dropdowns_are_filtered_by_gender(self):
        form = PersonForm()

        self.assertIn(self.male, form.fields["father"].queryset)
        self.assertNotIn(self.unknown, form.fields["father"].queryset)
        self.assertNotIn(self.female, form.fields["father"].queryset)
        self.assertIn(self.female, form.fields["mother"].queryset)
        self.assertNotIn(self.unknown, form.fields["mother"].queryset)
        self.assertNotIn(self.male, form.fields["mother"].queryset)

    def test_family_links_form_only_shows_father_and_mother(self):
        form = PersonForm()

        self.assertIn("father", form.fields)
        self.assertIn("mother", form.fields)
        self.assertNotIn("family_branch", form.fields)
        self.assertNotIn("household", form.fields)
        self.assertNotIn("generation", form.fields)
        self.assertNotIn("relationship_to_founder", form.fields)

    def test_member_relative_form_uses_same_parent_filters(self):
        form = MemberRelativeForm(show_workflow=False)

        self.assertNotIn(self.female, form.fields["father"].queryset)
        self.assertNotIn(self.male, form.fields["mother"].queryset)

    def test_invalid_parent_gender_is_rejected(self):
        form = PersonForm(
            data={
                "member_id": "",
                "first_name": "Child",
                "middle_name": "",
                "last_name": "Moshi",
                "clan_name": "Moshi",
                "gender": Person.Gender.MALE,
                "is_living": "on",
                "father": self.female.pk,
                "mother": self.male.pk,
                "status": Status.VERIFIED,
            }
        )

        self.assertFalse(form.is_valid())
        self.assertIn("father", form.errors)
        self.assertIn("mother", form.errors)


class ResidenceChoiceTests(TestCase):
    def test_residence_field_lists_mainland_zanzibar_and_outside_tanzania(self):
        form = PersonForm()
        rendered_choices = str(form["current_residence"])

        self.assertIn("Dar es Salaam", rendered_choices)
        self.assertIn(">Zanzibar<", rendered_choices)
        self.assertNotIn("Mjini Magharibi", rendered_choices)
        self.assertIn("Other (Nje ya Tanzania)", rendered_choices)

    def test_tanzania_locations_assign_country_automatically(self):
        form = PersonForm(
            data={
                "first_name": "Tanzania",
                "last_name": "Member",
                "clan_name": "Moshi",
                "gender": Person.Gender.MALE,
                "current_residence": "Arusha",
                "country": "",
                "is_living": "on",
                "status": Status.VERIFIED,
            }
        )

        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["country"], "Tanzania")

    def test_outside_tanzania_location_requires_country(self):
        form = PersonForm(
            data={
                "first_name": "Outside",
                "last_name": "Member",
                "clan_name": "Moshi",
                "gender": Person.Gender.FEMALE,
                "current_residence": "Other (Nje ya Tanzania)",
                "country": "",
                "is_living": "on",
                "status": Status.VERIFIED,
            }
        )

        self.assertFalse(form.is_valid())
        self.assertIn("country", form.errors)

    def test_existing_nonstandard_residence_is_preserved_for_editing(self):
        person = Person.objects.create(
            member_id="MOSHI-RESIDENCE-1",
            first_name="Existing",
            last_name="Member",
            current_residence="Nairobi",
        )

        form = PersonForm(instance=person)

        self.assertIn(("Nairobi", "Nairobi (Existing value)"), form.fields["current_residence"].choices)


class LivingStatusControlTests(TestCase):
    def test_date_fields_use_today_as_the_latest_selectable_date(self):
        form = PersonForm()
        today = timezone.localdate().isoformat()

        self.assertEqual(form.fields["date_of_birth"].widget.attrs["max"], today)
        self.assertEqual(form.fields["date_of_death"].widget.attrs["max"], today)
        self.assertIn("date-picker", form.fields["date_of_birth"].widget.attrs["class"])
        self.assertEqual(form.fields["date_of_birth"].widget.input_type, "text")
        self.assertEqual(form.fields["date_of_birth"].widget.attrs["autocomplete"], "off")

    def test_future_birth_and_death_dates_are_rejected(self):
        tomorrow = timezone.localdate() + timedelta(days=1)
        form = PersonForm(
            data={
                "first_name": "Future",
                "last_name": "Moshi",
                "gender": Person.Gender.MALE,
                "is_living": "",
                "date_of_birth": tomorrow.isoformat(),
                "date_of_death": tomorrow.isoformat(),
                "status": Status.VERIFIED,
            }
        )

        self.assertFalse(form.is_valid())
        self.assertIn("date_of_birth", form.errors)
        self.assertIn("date_of_death", form.errors)

    def test_date_of_death_cannot_precede_date_of_birth(self):
        form = PersonForm(
            data={
                "first_name": "Timeline",
                "last_name": "Moshi",
                "gender": Person.Gender.FEMALE,
                "is_living": "",
                "date_of_birth": "2000-01-01",
                "date_of_death": "1999-12-31",
                "status": Status.VERIFIED,
            }
        )

        self.assertFalse(form.is_valid())
        self.assertIn("date_of_death", form.errors)

    def test_living_person_form_clears_death_details(self):
        form = PersonForm(
            data={
                "member_id": "",
                "first_name": "Living",
                "middle_name": "",
                "last_name": "Moshi",
                "clan_name": "Moshi",
                "gender": Person.Gender.MALE,
                "is_living": "on",
                "date_of_death": "2020-01-01",
                "place_of_death": "Old value",
                "burial_location": "Old burial",
                "memorial_information": "Old memorial",
                "status": Status.VERIFIED,
            }
        )

        self.assertTrue(form.is_valid(), form.errors)
        person = form.save()
        self.assertIsNone(person.date_of_death)
        self.assertEqual(person.place_of_death, "")
        self.assertEqual(person.burial_location, "")
        self.assertEqual(person.memorial_information, "")

    def test_deceased_person_detail_shows_death_details(self):
        person = Person.objects.create(
            member_id="MOSHI-200",
            first_name="Legacy",
            last_name="Moshi",
            is_living=False,
            date_of_death="2020-01-01",
            place_of_death="Moshi",
            burial_location="Family cemetery",
            memorial_information="Remembered by the clan.",
        )

        response = self.client.get(reverse("person-detail", kwargs={"pk": person.pk}))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Deceased")
        self.assertContains(response, "Date of Death")
        self.assertContains(response, "Family cemetery")

    def test_living_person_detail_hides_death_labels(self):
        person = Person.objects.create(
            member_id="MOSHI-201",
            first_name="Alive",
            last_name="Moshi",
            is_living=True,
        )

        response = self.client.get(reverse("person-detail", kwargs={"pk": person.pk}))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Living")
        self.assertNotContains(response, "Date of Death")
